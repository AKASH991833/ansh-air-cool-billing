"""
Enterprise-Grade Profit/Loss, GST, Parts Usage & Team Performance Reports
Features 100% accurate financial calculations, multi-period filtering (Yearly, Quarterly, Monthly, Weekly),
clean enterprise white/slate styling, bold table totals, and Data-Analyst grade Excel export.
"""
from PySide6.QtWidgets import (
    QWidget, QVBoxLayout, QHBoxLayout, QFrame, QLabel, QPushButton,
    QComboBox, QTableWidget, QTableWidgetItem, QHeaderView, QScrollArea,
    QSizePolicy, QTabWidget, QGridLayout, QFileDialog, QMenu, QButtonGroup,
    QProgressBar, QDialog
)
from PySide6.QtCore import Qt, QTimer, Signal
from PySide6.QtGui import QFont, QColor, QCursor, QPixmap, QBrush
from datetime import datetime, date, timedelta
from decimal import Decimal
import os

from utils.unified_theme import UnifiedTheme
from views.base_window import BaseView, MetricCard
from utils.app_settings import get_setting

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
CHEVRON_ICON_PATH = os.path.join(BASE_DIR, "assets", "icons", "chevron_down.png").replace("\\", "/")


class ReportView(BaseView):
    def __init__(self):
        super().__init__()
        self.current_period = 'yearly'  # 'yearly', 'quarterly', 'monthly', 'weekly'
        self.cached_results = {}
        self._setup_ui()
        QTimer.singleShot(100, self.load_data)

    def update_theme_colors(self):
        colors = self.theme_manager.get_colors()
        self.theme_manager.apply_palette(self)
        self.setStyleSheet(self.theme_manager.get_main_stylesheet())

    def _make_pill_filter(self, text, checked=False):
        colors = self.theme_manager.get_colors()
        btn = QPushButton(text)
        btn.setCheckable(True)
        btn.setChecked(checked)
        btn.setCursor(QCursor(Qt.CursorShape.PointingHandCursor))
        btn.setStyleSheet(f"""
            QPushButton {{
                background: {colors['card_bg']};
                color: {colors['muted']};
                border: 1px solid {colors['border']};
                border-radius: 16px;
                padding: 6px 16px;
                font-size: 9pt;
                font-weight: 600;
            }}
            QPushButton:hover {{
                background: {colors['hover']};
                color: {colors['fg']};
                border-color: {colors['primary']}80;
            }}
            QPushButton:checked {{
                background: qlineargradient(x1:0, y1:0, x2:1, y2:1,
                    stop:0 {colors['primary']},
                    stop:1 #0891b2);
                color: #ffffff;
                border: 1px solid {colors['primary']};
                font-weight: 700;
            }}
        """)
        return btn

    def _make_combo(self, width=100):
        colors = self.theme_manager.get_colors()
        combo = QComboBox()
        combo.setMinimumWidth(width)
        combo.setStyleSheet(f"""
            QComboBox {{
                background: {colors['card_bg']};
                color: {colors['fg']};
                border: 1px solid {colors['border']};
                border-radius: 8px;
                padding: 6px 12px;
                font-size: 9pt;
                font-weight: 600;
                min-height: 20px;
            }}
            QComboBox:hover {{
                border-color: {colors['primary']};
            }}
            QComboBox::drop-down {{
                subcontrol-origin: padding;
                subcontrol-position: top right;
                border: none;
                width: 26px;
            }}
            QComboBox::down-arrow {{
                image: url("{CHEVRON_ICON_PATH}");
                width: 12px;
                height: 12px;
                margin-right: 6px;
            }}
            QComboBox QAbstractItemView {{
                background-color: {colors['card_bg']};
                color: {colors['fg']};
                border: 1px solid {colors['border']};
                border-radius: 8px;
                padding: 4px;
                selection-background-color: {colors['primary']};
                selection-color: #ffffff;
                outline: none;
            }}
            QComboBox QAbstractItemView::item {{
                padding: 6px 12px;
                border-radius: 4px;
            }}
        """)
        return combo

    def _setup_ui(self):
        main_layout = QVBoxLayout(self)
        main_layout.setContentsMargins(24, 20, 24, 20)
        main_layout.setSpacing(16)

        colors = self.theme_manager.get_colors()
        self.currency_symbol = get_setting("currency_symbol", "₹")

        # ── Header ──
        header_layout = QHBoxLayout()
        header_layout.setSpacing(16)

        title_frame = QFrame()
        title_layout = QVBoxLayout(title_frame)
        title_layout.setContentsMargins(0, 0, 0, 0)
        title_layout.setSpacing(2)

        title_row = QHBoxLayout()
        title_row.setSpacing(10)
        title_icon = QLabel("📊")
        title_icon.setStyleSheet("font-size: 22pt; background: transparent;")
        title_row.addWidget(title_icon)

        title_texts = QVBoxLayout()
        title_texts.setSpacing(0)
        title_lbl = QLabel("Reports & Analytics")
        title_lbl.setStyleSheet(f"font-size: 20pt; font-weight: 800; color: {colors['fg']}; letter-spacing: -0.5px;")
        title_texts.addWidget(title_lbl)
        subtitle_lbl = QLabel("Executive Financial, GST, Inventory & Workforce Intelligence")
        subtitle_lbl.setStyleSheet(f"font-size: 9pt; color: {colors['muted']}; font-weight: 500;")
        title_texts.addWidget(subtitle_lbl)
        title_row.addLayout(title_texts)
        title_layout.addLayout(title_row)

        header_layout.addWidget(title_frame)
        header_layout.addStretch()

        # ── Multi-Period Filter Bar ──
        filter_frame = QFrame()
        filter_frame.setObjectName("filterBar")
        filter_frame.setStyleSheet(f"""
            QFrame#filterBar {{
                background-color: {colors['card_bg']};
                border: 1px solid {colors['border']};
                border-radius: 12px;
                padding: 4px 8px;
            }}
        """)
        filter_layout = QHBoxLayout(filter_frame)
        filter_layout.setContentsMargins(6, 4, 6, 4)
        filter_layout.setSpacing(6)

        self.filter_group = QButtonGroup(self)
        periods = [
            ('yearly', '📅 Yearly'),
            ('quarterly', '📆 Quarterly'),
            ('monthly', '📋 Monthly'),
            ('weekly', '⏱️ Weekly'),
        ]
        for i, (key, label) in enumerate(periods):
            btn = self._make_pill_filter(label, checked=(key == 'yearly'))
            self.filter_group.addButton(btn, i)
            filter_layout.addWidget(btn)
            btn.clicked.connect(lambda checked, k=key: self._on_period_changed(k))

        sep = QFrame()
        sep.setFixedWidth(1)
        sep.setFixedHeight(22)
        sep.setStyleSheet(f"background: {colors['border']};")
        filter_layout.addWidget(sep)

        # Year Dropdown
        self.year_combo = self._make_combo(85)
        cy = datetime.now().year
        # Include years around 2026 and current year
        years = sorted(list(set([cy - 2, cy - 1, cy, cy + 1, 2025, 2026])), reverse=True)
        for y in years:
            self.year_combo.addItem(str(y))
        # Default to 2026 if 2026 is present, else cy
        def_year = "2026" if "2026" in [str(y) for y in years] else str(cy)
        self.year_combo.setCurrentText(def_year)
        self.year_combo.currentTextChanged.connect(lambda: self.load_data())
        filter_layout.addWidget(self.year_combo)

        # Quarter Dropdown (Visible on Quarterly)
        self.quarter_combo = self._make_combo(120)
        self.quarter_combo.addItem("Q1 (Jan - Mar)", 1)
        self.quarter_combo.addItem("Q2 (Apr - Jun)", 2)
        self.quarter_combo.addItem("Q3 (Jul - Sep)", 3)
        self.quarter_combo.addItem("Q4 (Oct - Dec)", 4)
        # Default to Q2 since May-June data is common
        self.quarter_combo.setCurrentIndex(1)
        self.quarter_combo.setVisible(False)
        self.quarter_combo.currentIndexChanged.connect(lambda: self.load_data())
        filter_layout.addWidget(self.quarter_combo)

        # Month Dropdown (Visible on Monthly)
        self.month_combo = self._make_combo(100)
        months = ['All Months', 'Jan', 'Feb', 'Mar', 'Apr', 'May', 'Jun', 'Jul', 'Aug', 'Sep', 'Oct', 'Nov', 'Dec']
        for m in months:
            self.month_combo.addItem(m)
        self.month_combo.setCurrentText('May')  # Useful default for current dataset
        self.month_combo.setVisible(False)
        self.month_combo.currentTextChanged.connect(lambda: self.load_data())
        filter_layout.addWidget(self.month_combo)

        # Week Range Dropdown (Visible on Weekly)
        self.week_combo = self._make_combo(125)
        self.week_combo.addItem("This Week", "this_week")
        self.week_combo.addItem("Last Week", "last_week")
        self.week_combo.addItem("Last 14 Days", "last_14")
        self.week_combo.addItem("Last 30 Days", "last_30")
        self.week_combo.addItem("May Peak (15-21 May)", "may_peak")
        self.week_combo.setVisible(False)
        self.week_combo.currentIndexChanged.connect(lambda: self.load_data())
        filter_layout.addWidget(self.week_combo)

        header_layout.addWidget(filter_frame)

        # ── Export Menu Button ──
        export_btn = QPushButton("  Export ▾")
        export_btn.setCursor(QCursor(Qt.CursorShape.PointingHandCursor))
        export_btn.setStyleSheet(f"""
            QPushButton {{
                background: qlineargradient(x1:0, y1:0, x2:1, y2:1,
                    stop:0 {colors['primary']},
                    stop:1 #0891b2);
                color: #ffffff;
                border: none;
                border-radius: 10px;
                padding: 8px 18px;
                font-size: 9.5pt;
                font-weight: 700;
            }}
            QPushButton:hover {{
                background: qlineargradient(x1:0, y1:0, x2:1, y2:1,
                    stop:0 #38bdf8,
                    stop:1 #0284c7);
            }}
        """)
        export_menu = QMenu(self)
        export_menu.setStyleSheet(f"""
            QMenu {{
                background-color: {colors['card_bg']};
                border: 1px solid {colors['border']};
                border-radius: 8px;
                padding: 6px;
            }}
            QMenu::item {{
                padding: 8px 24px;
                border-radius: 6px;
                font-size: 9pt;
                color: {colors['fg']};
            }}
            QMenu::item:selected {{
                background: {colors['primary']}25;
                color: {colors['primary']};
            }}
            QMenu::separator {{
                height: 1px;
                background: {colors['border']};
                margin: 4px 10px;
            }}
        """)
        export_menu.addAction("📊  Export Active Tab to Excel", lambda: self._export_current_tab('excel'))
        export_menu.addAction("📑  Export Complete Executive Workbook (All Tabs)", lambda: self._export_master_workbook())
        export_menu.addSeparator()
        export_menu.addAction("📄  Export to PDF", lambda: self._export_current_tab('pdf'))
        export_menu.addAction("🖨️  Print Report", lambda: self._export_current_tab('print'))
        export_btn.setMenu(export_menu)
        header_layout.addWidget(export_btn)

        main_layout.addLayout(header_layout)

        # ── 8 Executive Metric KPI Cards ──
        self.summary_frame = QFrame()
        self.summary_frame.setObjectName("metricContainer")
        self.summary_frame.setStyleSheet(f"""
            QFrame#metricContainer {{
                background: {colors['card_bg']};
                border: 1px solid {colors['border']};
                border-radius: 14px;
                padding: 12px;
            }}
        """)
        summary_grid = QGridLayout(self.summary_frame)
        summary_grid.setContentsMargins(4, 4, 4, 4)
        summary_grid.setSpacing(10)

        self.cards = {}
        metrics = [
            ('gross_revenue', 'Gross Revenue', f'{self.currency_symbol}0', '#0284c7', 'Total invoiced revenue across services & parts'),
            ('service_revenue', 'Service Revenue', f'{self.currency_symbol}0', '#2563eb', 'Total revenue generated from labor & services'),
            ('part_revenue', 'Parts Revenue', f'{self.currency_symbol}0', '#7c3aed', 'Total sales value of replacement parts & hardware'),
            ('net_profit', 'Net Profit', f'{self.currency_symbol}0', '#059669', 'Net business profit after parts, petrol & commission'),
            ('total_expenses', 'Total Expenses', f'{self.currency_symbol}0', '#dc2626', 'Sum of parts cost, petrol expenses & commissions'),
            ('total_invoices', 'Total Invoices', '0', '#d97706', 'Total count of official invoices generated'),
            ('total_visits', 'Field Visits', '0', '#475569', 'Total field visits recorded in Daily Logs & Invoices'),
            ('collection_rate', 'Collection Rate', '0.0%', '#0d9488', 'Percentage of invoiced billing successfully collected'),
        ]
        for i, (key, label, default, color, tooltip) in enumerate(metrics):
            card = MetricCard(label, default, self._get_icon(key), color)
            card.setToolTip(tooltip)
            self.cards[key] = card
            summary_grid.addWidget(card, i // 4, i % 4)

        main_layout.addWidget(self.summary_frame)

        # ── Enterprise Tabbed Data Tables (CHARTS REMOVED COMPLETELY) ──
        self.content_frame = QFrame()
        self.content_frame.setStyleSheet(f"""
            QFrame {{
                background: {colors['card_bg']};
                border: 1px solid {colors['border']};
                border-radius: 14px;
            }}
        """)
        content_layout = QVBoxLayout(self.content_frame)
        content_layout.setContentsMargins(14, 12, 14, 12)
        content_layout.setSpacing(10)

        self.tabs = QTabWidget()
        self.tabs.setStyleSheet(f"""
            QTabWidget::pane {{
                border: none;
                background: transparent;
                margin-top: 4px;
            }}
            QTabBar::tab {{
                background: transparent;
                color: {colors['muted']};
                padding: 10px 20px;
                font-weight: 700;
                font-size: 9.5pt;
                border: none;
                border-bottom: 3px solid transparent;
                margin-right: 6px;
            }}
            QTabBar::tab:selected {{
                color: {colors['primary']};
                border-bottom: 3px solid {colors['primary']};
            }}
            QTabBar::tab:hover:!selected {{
                color: {colors['fg']};
                background: {colors['hover']}60;
                border-radius: 6px 6px 0 0;
            }}
        """)

        self._build_pnl_tab()
        self._build_gst_tab()
        self._build_parts_tab()
        self._build_team_tab()

        content_layout.addWidget(self.tabs)
        main_layout.addWidget(self.content_frame, 1)

    def _get_icon(self, key):
        icons = {
            'gross_revenue': '💰',
            'service_revenue': '🔧',
            'part_revenue': '⚙️',
            'net_profit': '📈',
            'total_expenses': '💸',
            'total_invoices': '📄',
            'total_visits': '🚗',
            'collection_rate': '🎯',
        }
        return icons.get(key, '📊')

    def _on_period_changed(self, period_key):
        self.current_period = period_key
        self.year_combo.setVisible(period_key in ['yearly', 'quarterly', 'monthly'])
        self.quarter_combo.setVisible(period_key == 'quarterly')
        self.month_combo.setVisible(period_key == 'monthly')
        self.week_combo.setVisible(period_key == 'weekly')
        self.load_data()

    def _make_table(self, cols, headers):
        colors = self.theme_manager.get_colors()
        table = QTableWidget()
        table.setColumnCount(cols)
        table.setHorizontalHeaderLabels(headers)
        table.setSelectionBehavior(QTableWidget.SelectionBehavior.SelectRows)
        table.setEditTriggers(QTableWidget.EditTrigger.NoEditTriggers)
        table.setAlternatingRowColors(True)
        table.verticalHeader().setVisible(False)
        table.verticalHeader().setDefaultSectionSize(38)
        table.setShowGrid(True)
        table.setSortingEnabled(False)
        table.setStyleSheet(f"""
            QTableWidget {{
                background-color: #ffffff;
                color: #0f172a;
                border: 1px solid {colors['border']};
                border-radius: 8px;
                gridline-color: #e2e8f0;
                selection-background-color: #e0f2fe;
                selection-color: #0369a1;
                alternate-background-color: #f8fafc;
            }}
            QTableWidget::item {{
                padding: 6px 12px;
                border-bottom: 1px solid #e2e8f0;
                font-size: 9pt;
            }}
            QHeaderView::section {{
                background-color: #f1f5f9;
                color: #334155;
                padding: 8px 10px;
                border: none;
                border-bottom: 2px solid #cbd5e1;
                border-right: 1px solid #e2e8f0;
                font-weight: 700;
                font-size: 8.5pt;
                letter-spacing: 0.2px;
            }}
            QHeaderView::section:last {{
                border-right: none;
            }}
        """)
        header = table.horizontalHeader()
        header.setSectionResizeMode(0, QHeaderView.ResizeMode.Stretch)
        for i in range(1, cols):
            header.setSectionResizeMode(i, QHeaderView.ResizeMode.ResizeToContents)
        return table

    def _make_info_bar(self, icon, text):
        colors = self.theme_manager.get_colors()
        bar = QFrame()
        bar.setStyleSheet(f"""
            QFrame {{
                background-color: {colors['hover']};
                border: 1px solid {colors['border']};
                border-radius: 8px;
                padding: 6px 12px;
            }}
        """)
        layout = QHBoxLayout(bar)
        layout.setContentsMargins(4, 2, 4, 2)
        layout.setSpacing(8)
        icon_lbl = QLabel(icon)
        icon_lbl.setStyleSheet("font-size: 11pt; background: transparent;")
        layout.addWidget(icon_lbl)
        text_lbl = QLabel(text)
        text_lbl.setStyleSheet(f"font-size: 8.5pt; color: {colors['muted']}; font-weight: 500; background: transparent;")
        layout.addWidget(text_lbl)
        layout.addStretch()
        return bar

    def _build_pnl_tab(self):
        tab = QWidget()
        layout = QVBoxLayout(tab)
        layout.setContentsMargins(4, 8, 4, 4)
        layout.setSpacing(8)

        layout.addWidget(self._make_info_bar(
            "📈",
            "Statement of Profit & Loss: Revenue, Material/Parts Cost, Petrol Expense, Technician Commission & Net Profit Margin"
        ))

        self.pnl_table = self._make_table(11, [
            'Period', 'Gross Revenue', 'Service Rev', 'Parts Rev',
            'Parts Cost', 'Petrol Cost', 'Commission', 'Total Expenses',
            'Net Profit', 'Margin %', 'Collection %'
        ])
        layout.addWidget(self.pnl_table)
        self.tabs.addTab(tab, "📈 Profit & Loss Statement")

    def _build_gst_tab(self):
        tab = QWidget()
        layout = QVBoxLayout(tab)
        layout.setContentsMargins(4, 8, 4, 4)
        layout.setSpacing(8)

        layout.addWidget(self._make_info_bar(
            "🧾",
            "GST Compliance Ledger: Taxable Invoiced Value, CGST (9%), SGST (9%), Total Tax Collected & Invoices Count"
        ))

        self.gst_table = self._make_table(6, [
            'Period', 'Invoices', 'Taxable Value', 'CGST (9%)', 'SGST (9%)', 'Total GST Collected'
        ])
        layout.addWidget(self.gst_table)
        self.tabs.addTab(tab, "🧾 GST Report")

    def _build_parts_tab(self):
        tab = QWidget()
        layout = QVBoxLayout(tab)
        layout.setContentsMargins(4, 8, 4, 4)
        layout.setSpacing(8)

        layout.addWidget(self._make_info_bar(
            "🔧",
            "Parts & Component Consumption: Units consumed across service jobs with revenue, cost price and financial contribution"
        ))

        self.parts_table = self._make_table(6, [
            'Part Name', 'Category', 'Quantity Used', 'Total Billed (Sales)', 'Total Cost (Purchase)', '% of Parts Total'
        ])
        layout.addWidget(self.parts_table)
        self.tabs.addTab(tab, "🔧 Parts Usage")

    def _build_team_tab(self):
        tab = QWidget()
        layout = QVBoxLayout(tab)
        layout.setContentsMargins(4, 8, 4, 4)
        layout.setSpacing(8)

        layout.addWidget(self._make_info_bar(
            "👥",
            "Technician Workforce Performance: Field visits completed, customers serviced, gross billing, commission payout & profit"
        ))

        self.tech_table = self._make_table(7, [
            'Technician Name', 'Field Visits', 'Customers Serviced',
            'Total Billed', 'Commission Payout', 'Net Profit Generated', 'Avg Profit / Visit'
        ])
        layout.addWidget(self.tech_table)
        self.tabs.addTab(tab, "👥 Team Performance")

    def _get_active_date_range(self):
        """Calculate start_date and end_date based on active period controls"""
        try:
            year = int(self.year_combo.currentText())
        except Exception:
            year = 2026

        if self.current_period == 'yearly':
            start_date = f"{year}-01-01"
            end_date = f"{year + 1}-01-01"
            period_label = f"Year {year}"
            return start_date, end_date, period_label

        elif self.current_period == 'quarterly':
            q = self.quarter_combo.currentData() or 1
            quarter_ranges = {
                1: (f"{year}-01-01", f"{year}-04-01", f"Q1 {year}"),
                2: (f"{year}-04-01", f"{year}-07-01", f"Q2 {year}"),
                3: (f"{year}-07-01", f"{year}-10-01", f"Q3 {year}"),
                4: (f"{year}-10-01", f"{year + 1}-01-01", f"Q4 {year}"),
            }
            s, e, lbl = quarter_ranges.get(q, (f"{year}-01-01", f"{year + 1}-01-01", f"Q1 {year}"))
            return s, e, lbl

        elif self.current_period == 'monthly':
            m_text = self.month_combo.currentText()
            if m_text == 'All Months':
                return f"{year}-01-01", f"{year + 1}-01-01", f"All Months {year}"
            months = ['Jan', 'Feb', 'Mar', 'Apr', 'May', 'Jun', 'Jul', 'Aug', 'Sep', 'Oct', 'Nov', 'Dec']
            m_idx = months.index(m_text) + 1 if m_text in months else 5
            if m_idx == 12:
                return f"{year}-12-01", f"{year + 1}-01-01", f"{m_text} {year}"
            else:
                return f"{year}-{m_idx:02d}-01", f"{year}-{m_idx + 1:02d}-01", f"{m_text} {year}"

        elif self.current_period == 'weekly':
            preset = self.week_combo.currentData() or "this_week"
            today = datetime.now().date()
            if preset == "this_week":
                start_dt = today - timedelta(days=today.weekday())
                end_dt = start_dt + timedelta(days=7)
                return start_dt.strftime('%Y-%m-%d'), end_dt.strftime('%Y-%m-%d'), "This Week"
            elif preset == "last_week":
                start_dt = today - timedelta(days=today.weekday() + 7)
                end_dt = start_dt + timedelta(days=7)
                return start_dt.strftime('%Y-%m-%d'), end_dt.strftime('%Y-%m-%d'), "Last Week"
            elif preset == "last_14":
                start_dt = today - timedelta(days=14)
                end_dt = today + timedelta(days=1)
                return start_dt.strftime('%Y-%m-%d'), end_dt.strftime('%Y-%m-%d'), "Last 14 Days"
            elif preset == "last_30":
                start_dt = today - timedelta(days=30)
                end_dt = today + timedelta(days=1)
                return start_dt.strftime('%Y-%m-%d'), end_dt.strftime('%Y-%m-%d'), "Last 30 Days"
            elif preset == "may_peak":
                return "2026-05-15", "2026-05-22", "May Peak (15-21 May 2026)"

        return f"{year}-01-01", f"{year + 1}-01-01", f"Year {year}"

    def load_data(self):
        self.show_loading("Aggregating financial & operational data...")
        self.run_in_thread(
            self._load_report_data_thread,
            self._update_report_ui,
            self._on_report_error
        )

    def _on_report_error(self, err):
        self.hide_loading()
        self._default_error_handler(err)

    def _load_report_data_thread(self):
        from database.db_connection import DatabaseContext
        from controllers.report_controller import ReportController

        start_date, end_date, period_label = self._get_active_date_range()

        try:
            year = int(self.year_combo.currentText())
        except Exception:
            year = 2026

        with DatabaseContext() as db:
            rc = ReportController(db)

            # 1. Period Breakdown for P&L Table
            if self.current_period == 'yearly':
                breakdown = rc.get_yearly_summary(year)
            elif self.current_period == 'quarterly':
                q = self.quarter_combo.currentData() or 1
                breakdown = rc.get_quarterly_summary(year, q)
            elif self.current_period == 'monthly':
                m_text = self.month_combo.currentText()
                if m_text == 'All Months':
                    breakdown = rc.get_yearly_summary(year)
                else:
                    months = ['Jan', 'Feb', 'Mar', 'Apr', 'May', 'Jun', 'Jul', 'Aug', 'Sep', 'Oct', 'Nov', 'Dec']
                    m_idx = months.index(m_text) + 1 if m_text in months else 5
                    # Break down the single month into 4 weeks
                    breakdown = [
                        rc.get_period_analytics(f"{year}-{m_idx:02d}-01", f"{year}-{m_idx:02d}-08", label=f"Week 1 ({m_text} 01-07)"),
                        rc.get_period_analytics(f"{year}-{m_idx:02d}-08", f"{year}-{m_idx:02d}-15", label=f"Week 2 ({m_text} 08-14)"),
                        rc.get_period_analytics(f"{year}-{m_idx:02d}-15", f"{year}-{m_idx:02d}-22", label=f"Week 3 ({m_text} 15-21)"),
                        rc.get_period_analytics(f"{year}-{m_idx:02d}-22", end_date, label=f"Week 4 ({m_text} 22-End)"),
                    ]
            elif self.current_period == 'weekly':
                breakdown = rc.get_weekly_summary(start_date, (datetime.strptime(end_date, '%Y-%m-%d') - timedelta(days=1)).strftime('%Y-%m-%d'))
            else:
                breakdown = rc.get_yearly_summary(year)

            # 2. Overall Summary KPIs for the Selected Active Range
            overall = rc.get_period_analytics(start_date, end_date, label=period_label)

            # 3. GST Report
            gst_report = rc.get_gst_report(start_date, end_date)

            # 4. Parts Usage
            parts_usage = rc.get_parts_usage_combined(start_date, end_date)

            # 5. Technician Performance
            tech_perf = rc.get_technician_performance(start_date, end_date)

            return {
                'period_label': period_label,
                'start_date': start_date,
                'end_date': end_date,
                'overall': overall,
                'breakdown': breakdown,
                'gst_report': gst_report,
                'parts_usage': parts_usage,
                'tech_perf': tech_perf,
            }

    def _update_report_ui(self, result):
        try:
            self.cached_results = result
            self._update_kpi_cards(result['overall'])
            self._update_pnl_table(result['breakdown'])
            self._update_gst_table(result['gst_report'])
            self._update_parts_table(result['parts_usage'])
            self._update_tech_table(result['tech_perf'])
        except Exception as e:
            print(f"Error updating report UI: {e}")
            import traceback
            traceback.print_exc()
        finally:
            self.hide_loading()

    def _update_kpi_cards(self, ov):
        if not ov:
            return

        gross = ov.get('gross_revenue', 0.0)
        svc = ov.get('service_revenue', 0.0)
        parts = ov.get('part_revenue', 0.0)
        exp = ov.get('total_expenses', 0.0)
        profit = ov.get('net_profit', 0.0)
        invs = ov.get('total_invoices', 0)
        visits = ov.get('total_visits', 0)
        coll = ov.get('collection_rate', 0.0)

        self.cards['gross_revenue'].set_value(f"{self.currency_symbol}{gross:,.0f}")
        self.cards['service_revenue'].set_value(f"{self.currency_symbol}{svc:,.0f}")
        self.cards['part_revenue'].set_value(f"{self.currency_symbol}{parts:,.0f}")
        self.cards['total_expenses'].set_value(f"{self.currency_symbol}{exp:,.0f}")
        self.cards['net_profit'].set_value(f"{self.currency_symbol}{profit:,.0f}")
        self.cards['total_invoices'].set_value(f"{invs:,}")
        self.cards['total_visits'].set_value(f"{visits:,}")
        self.cards['collection_rate'].set_value(f"{coll:.1f}%")

        # Color profit card green or red
        self.cards['net_profit'].set_color('#059669' if profit >= 0 else '#dc2626')

    def _format_cell(self, text, align=Qt.AlignCenter, color=None, bold=False, is_total=False):
        item = QTableWidgetItem(str(text))
        item.setTextAlignment(align | Qt.AlignVCenter)
        if color:
            item.setForeground(QColor(color))
        if bold:
            item.setFont(QFont("Segoe UI", 9, QFont.Weight.Bold))
        if is_total:
            item.setBackground(QColor("#eff6ff"))
            item.setFont(QFont("Segoe UI", 9, QFont.Weight.Bold))
            if not color:
                item.setForeground(QColor("#0f172a"))
        return item

    def _update_pnl_table(self, breakdown):
        self.pnl_table.setRowCount(0)
        if not breakdown:
            return

        tot_rev = 0.0
        tot_svc = 0.0
        tot_parts_rev = 0.0
        tot_parts_cost = 0.0
        tot_petrol = 0.0
        tot_comm = 0.0
        tot_exp = 0.0
        tot_profit = 0.0
        tot_collected = 0.0

        for d in breakdown:
            row = self.pnl_table.rowCount()
            self.pnl_table.insertRow(row)

            rev = d.get('gross_revenue', 0.0)
            svc = d.get('service_revenue', 0.0)
            prev = d.get('part_revenue', 0.0)
            pcost = d.get('parts_cost', 0.0)
            petrol = d.get('petrol_cost', 0.0)
            comm = d.get('commission', 0.0)
            exp = d.get('total_expenses', 0.0)
            pf = d.get('net_profit', 0.0)
            margin = d.get('margin_pct', 0.0)
            coll = d.get('collection_rate', 0.0)

            tot_rev += rev
            tot_svc += svc
            tot_parts_rev += prev
            tot_parts_cost += pcost
            tot_petrol += petrol
            tot_comm += comm
            tot_exp += exp
            tot_profit += pf
            tot_collected += d.get('total_collected', 0.0)

            self.pnl_table.setItem(row, 0, self._format_cell(d.get('label', ''), Qt.AlignLeft, bold=True))
            self.pnl_table.setItem(row, 1, self._format_cell(f"{self.currency_symbol}{rev:,.0f}", Qt.AlignRight, '#0284c7'))
            self.pnl_table.setItem(row, 2, self._format_cell(f"{self.currency_symbol}{svc:,.0f}", Qt.AlignRight))
            self.pnl_table.setItem(row, 3, self._format_cell(f"{self.currency_symbol}{prev:,.0f}", Qt.AlignRight))
            self.pnl_table.setItem(row, 4, self._format_cell(f"{self.currency_symbol}{pcost:,.0f}", Qt.AlignRight, '#7c3aed'))
            self.pnl_table.setItem(row, 5, self._format_cell(f"{self.currency_symbol}{petrol:,.0f}", Qt.AlignRight, '#d97706'))
            self.pnl_table.setItem(row, 6, self._format_cell(f"{self.currency_symbol}{comm:,.0f}", Qt.AlignRight, '#475569'))
            self.pnl_table.setItem(row, 7, self._format_cell(f"{self.currency_symbol}{exp:,.0f}", Qt.AlignRight, '#dc2626'))

            pf_color = '#059669' if pf >= 0 else '#dc2626'
            self.pnl_table.setItem(row, 8, self._format_cell(f"{self.currency_symbol}{pf:,.0f}", Qt.AlignRight, pf_color, bold=True))
            self.pnl_table.setItem(row, 9, self._format_cell(f"{margin:.1f}%", Qt.AlignCenter, pf_color, bold=True))
            self.pnl_table.setItem(row, 10, self._format_cell(f"{coll:.1f}%", Qt.AlignCenter, '#0d9488'))

        # Add Bold Total Summary Row
        tot_row = self.pnl_table.rowCount()
        self.pnl_table.insertRow(tot_row)
        overall_margin = (tot_profit / tot_rev * 100.0) if tot_rev > 0 else 0.0
        overall_coll = (tot_collected / tot_rev * 100.0) if tot_rev > 0 else 0.0

        self.pnl_table.setItem(tot_row, 0, self._format_cell("TOTAL SUMMARY", Qt.AlignLeft, bold=True, is_total=True))
        self.pnl_table.setItem(tot_row, 1, self._format_cell(f"{self.currency_symbol}{tot_rev:,.0f}", Qt.AlignRight, '#0284c7', bold=True, is_total=True))
        self.pnl_table.setItem(tot_row, 2, self._format_cell(f"{self.currency_symbol}{tot_svc:,.0f}", Qt.AlignRight, bold=True, is_total=True))
        self.pnl_table.setItem(tot_row, 3, self._format_cell(f"{self.currency_symbol}{tot_parts_rev:,.0f}", Qt.AlignRight, bold=True, is_total=True))
        self.pnl_table.setItem(tot_row, 4, self._format_cell(f"{self.currency_symbol}{tot_parts_cost:,.0f}", Qt.AlignRight, '#7c3aed', bold=True, is_total=True))
        self.pnl_table.setItem(tot_row, 5, self._format_cell(f"{self.currency_symbol}{tot_petrol:,.0f}", Qt.AlignRight, '#d97706', bold=True, is_total=True))
        self.pnl_table.setItem(tot_row, 6, self._format_cell(f"{self.currency_symbol}{tot_comm:,.0f}", Qt.AlignRight, '#475569', bold=True, is_total=True))
        self.pnl_table.setItem(tot_row, 7, self._format_cell(f"{self.currency_symbol}{tot_exp:,.0f}", Qt.AlignRight, '#dc2626', bold=True, is_total=True))

        tot_pf_color = '#059669' if tot_profit >= 0 else '#dc2626'
        self.pnl_table.setItem(tot_row, 8, self._format_cell(f"{self.currency_symbol}{tot_profit:,.0f}", Qt.AlignRight, tot_pf_color, bold=True, is_total=True))
        self.pnl_table.setItem(tot_row, 9, self._format_cell(f"{overall_margin:.1f}%", Qt.AlignCenter, tot_pf_color, bold=True, is_total=True))
        self.pnl_table.setItem(tot_row, 10, self._format_cell(f"{overall_coll:.1f}%", Qt.AlignCenter, '#0d9488', bold=True, is_total=True))

    def _update_gst_table(self, gst_rows):
        self.gst_table.setRowCount(0)
        if not gst_rows:
            return

        tot_inv = 0
        tot_taxable = 0.0
        tot_cgst = 0.0
        tot_sgst = 0.0
        tot_gst = 0.0

        for r in gst_rows:
            row = self.gst_table.rowCount()
            self.gst_table.insertRow(row)

            inv_count = r.get('invoice_count', 0)
            taxable = r.get('taxable_value', 0.0)
            cgst = r.get('cgst', 0.0)
            sgst = r.get('sgst', 0.0)
            gst = r.get('total_gst', 0.0)

            tot_inv += inv_count
            tot_taxable += taxable
            tot_cgst += cgst
            tot_sgst += sgst
            tot_gst += gst

            self.gst_table.setItem(row, 0, self._format_cell(r.get('month', ''), Qt.AlignCenter, bold=True))
            self.gst_table.setItem(row, 1, self._format_cell(f"{inv_count:,}", Qt.AlignCenter))
            self.gst_table.setItem(row, 2, self._format_cell(f"{self.currency_symbol}{taxable:,.0f}", Qt.AlignRight))
            self.gst_table.setItem(row, 3, self._format_cell(f"{self.currency_symbol}{cgst:,.0f}", Qt.AlignRight))
            self.gst_table.setItem(row, 4, self._format_cell(f"{self.currency_symbol}{sgst:,.0f}", Qt.AlignRight))
            self.gst_table.setItem(row, 5, self._format_cell(f"{self.currency_symbol}{gst:,.0f}", Qt.AlignRight, '#0284c7', bold=True))

        # Add Total GST Summary Row
        tot_row = self.gst_table.rowCount()
        self.gst_table.insertRow(tot_row)
        self.gst_table.setItem(tot_row, 0, self._format_cell("TOTAL GST", Qt.AlignCenter, bold=True, is_total=True))
        self.gst_table.setItem(tot_row, 1, self._format_cell(f"{tot_inv:,}", Qt.AlignCenter, bold=True, is_total=True))
        self.gst_table.setItem(tot_row, 2, self._format_cell(f"{self.currency_symbol}{tot_taxable:,.0f}", Qt.AlignRight, bold=True, is_total=True))
        self.gst_table.setItem(tot_row, 3, self._format_cell(f"{self.currency_symbol}{tot_cgst:,.0f}", Qt.AlignRight, bold=True, is_total=True))
        self.gst_table.setItem(tot_row, 4, self._format_cell(f"{self.currency_symbol}{tot_sgst:,.0f}", Qt.AlignRight, bold=True, is_total=True))
        self.gst_table.setItem(tot_row, 5, self._format_cell(f"{self.currency_symbol}{tot_gst:,.0f}", Qt.AlignRight, '#0284c7', bold=True, is_total=True))

    def _update_parts_table(self, parts):
        self.parts_table.setRowCount(0)
        if not parts:
            return

        total_parts_rev = sum(p.get('total_revenue', 0.0) for p in parts) or 1.0
        tot_qty = 0
        tot_rev = 0.0
        tot_cost = 0.0

        for r in parts:
            row = self.parts_table.rowCount()
            self.parts_table.insertRow(row)

            qty = r.get('qty_used', 0)
            rev = r.get('total_revenue', 0.0)
            cost = r.get('total_cost', 0.0)
            pct = (rev / total_parts_rev * 100.0)

            tot_qty += qty
            tot_rev += rev
            tot_cost += cost

            self.parts_table.setItem(row, 0, self._format_cell(r.get('part_name', ''), Qt.AlignLeft, bold=True))
            self.parts_table.setItem(row, 1, self._format_cell(r.get('category', 'General'), Qt.AlignCenter, '#64748b'))
            self.parts_table.setItem(row, 2, self._format_cell(f"{qty:,}", Qt.AlignCenter))
            self.parts_table.setItem(row, 3, self._format_cell(f"{self.currency_symbol}{rev:,.0f}", Qt.AlignRight, '#0284c7'))
            self.parts_table.setItem(row, 4, self._format_cell(f"{self.currency_symbol}{cost:,.0f}", Qt.AlignRight, '#7c3aed'))
            self.parts_table.setItem(row, 5, self._format_cell(f"{pct:.1f}%", Qt.AlignCenter, '#059669', bold=True))

        # Add Total Summary Row
        tot_row = self.parts_table.rowCount()
        self.parts_table.insertRow(tot_row)
        self.parts_table.setItem(tot_row, 0, self._format_cell("TOTAL CONSUMPTION", Qt.AlignLeft, bold=True, is_total=True))
        self.parts_table.setItem(tot_row, 1, self._format_cell(f"{len(parts)} Items", Qt.AlignCenter, bold=True, is_total=True))
        self.parts_table.setItem(tot_row, 2, self._format_cell(f"{tot_qty:,}", Qt.AlignCenter, bold=True, is_total=True))
        self.parts_table.setItem(tot_row, 3, self._format_cell(f"{self.currency_symbol}{tot_rev:,.0f}", Qt.AlignRight, '#0284c7', bold=True, is_total=True))
        self.parts_table.setItem(tot_row, 4, self._format_cell(f"{self.currency_symbol}{tot_cost:,.0f}", Qt.AlignRight, '#7c3aed', bold=True, is_total=True))
        self.parts_table.setItem(tot_row, 5, self._format_cell("100.0%", Qt.AlignCenter, '#059669', bold=True, is_total=True))

    def _update_tech_table(self, tech_data):
        self.tech_table.setRowCount(0)
        if not tech_data:
            return

        tot_visits = 0
        tot_cust = 0
        tot_billed = 0.0
        tot_comm = 0.0
        tot_profit = 0.0

        for d in tech_data:
            row = self.tech_table.rowCount()
            self.tech_table.insertRow(row)

            visits = d.get('total_visits', 0)
            custs = d.get('unique_customers', 0)
            pay = d.get('total_payment', 0.0)
            exp = d.get('total_expense', 0.0)
            pf = d.get('total_profit', 0.0)
            avg = d.get('avg_profit_per_visit', 0.0)

            tot_visits += visits
            tot_cust += custs
            tot_billed += pay
            tot_comm += exp
            tot_profit += pf

            pf_color = '#059669' if pf >= 0 else '#dc2626'

            self.tech_table.setItem(row, 0, self._format_cell(d.get('technician_name', ''), Qt.AlignLeft, bold=True))
            self.tech_table.setItem(row, 1, self._format_cell(f"{visits:,}", Qt.AlignCenter))
            self.tech_table.setItem(row, 2, self._format_cell(f"{custs:,}", Qt.AlignCenter))
            self.tech_table.setItem(row, 3, self._format_cell(f"{self.currency_symbol}{pay:,.0f}", Qt.AlignRight, '#0284c7'))
            self.tech_table.setItem(row, 4, self._format_cell(f"{self.currency_symbol}{exp:,.0f}", Qt.AlignRight, '#dc2626'))
            self.tech_table.setItem(row, 5, self._format_cell(f"{self.currency_symbol}{pf:,.0f}", Qt.AlignRight, pf_color, bold=True))
            self.tech_table.setItem(row, 6, self._format_cell(f"{self.currency_symbol}{avg:,.0f}", Qt.AlignRight, '#0d9488'))

        # Add Total Summary Row
        tot_row = self.tech_table.rowCount()
        self.tech_table.insertRow(tot_row)
        overall_avg = (tot_profit / tot_visits) if tot_visits > 0 else 0.0
        tot_pf_color = '#059669' if tot_profit >= 0 else '#dc2626'

        self.tech_table.setItem(tot_row, 0, self._format_cell("TOTAL TEAM", Qt.AlignLeft, bold=True, is_total=True))
        self.tech_table.setItem(tot_row, 1, self._format_cell(f"{tot_visits:,}", Qt.AlignCenter, bold=True, is_total=True))
        self.tech_table.setItem(tot_row, 2, self._format_cell(f"{tot_cust:,}", Qt.AlignCenter, bold=True, is_total=True))
        self.tech_table.setItem(tot_row, 3, self._format_cell(f"{self.currency_symbol}{tot_billed:,.0f}", Qt.AlignRight, '#0284c7', bold=True, is_total=True))
        self.tech_table.setItem(tot_row, 4, self._format_cell(f"{self.currency_symbol}{tot_comm:,.0f}", Qt.AlignRight, '#dc2626', bold=True, is_total=True))
        self.tech_table.setItem(tot_row, 5, self._format_cell(f"{self.currency_symbol}{tot_profit:,.0f}", Qt.AlignRight, tot_pf_color, bold=True, is_total=True))
        self.tech_table.setItem(tot_row, 6, self._format_cell(f"{self.currency_symbol}{overall_avg:,.0f}", Qt.AlignRight, '#0d9488', bold=True, is_total=True))

    def _export_current_tab(self, fmt='excel'):
        """Export the currently active tab with analyst grade formatting"""
        from utils.excel_helper import ExcelExporter

        tab_idx = self.tabs.currentIndex()
        tab_names = {0: 'Profit & Loss', 1: 'GST Report', 2: 'Parts Usage', 3: 'Team Performance'}
        table_map = {0: self.pnl_table, 1: self.gst_table, 2: self.parts_table, 3: self.tech_table}
        table = table_map.get(tab_idx)
        tab_name = tab_names.get(tab_idx, 'Report')

        if not table or table.rowCount() == 0:
            self.show_warning_message("No data in this tab to export")
            return

        headers = [table.horizontalHeaderItem(i).text() for i in range(table.columnCount())]
        rows = []
        for r in range(table.rowCount()):
            row_data = []
            for c in range(table.columnCount()):
                item = table.item(r, c)
                row_data.append(item.text() if item else '')
            rows.append(row_data)

        period_lbl = self.cached_results.get('period_label', 'Custom Period')

        if fmt == 'excel':
            kpis = [
                ('Report', tab_name),
                ('Period', period_lbl),
                ('Total Rows', len(rows)),
            ]
            ov = self.cached_results.get('overall', {})
            if ov:
                kpis.append(('Gross Revenue', f"{self.currency_symbol}{ov.get('gross_revenue', 0):,.0f}"))
                kpis.append(('Net Profit', f"{self.currency_symbol}{ov.get('net_profit', 0):,.0f}"))

            wb = ExcelExporter.build_excel(
                sheet_title=tab_name[:30],
                title_text=f'{tab_name} Statement',
                subtitle_text=f'Period: {period_lbl} | Generated: {datetime.now().strftime("%d-%m-%Y %H:%M")}',
                headers=headers,
                rows=rows,
                kpis=kpis
            )
            safe_name = tab_name.replace(' ', '_').lower()
            safe_period = period_lbl.replace(' ', '_').replace('(', '').replace(')', '').lower()
            filepath, _ = QFileDialog.getSaveFileName(
                self, 'Save Analyst Excel Report',
                f"{safe_name}_{safe_period}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx",
                'Excel Files (*.xlsx)'
            )
            if filepath:
                try:
                    wb.save(filepath)
                    self.show_success_message(f"Report exported successfully:\n{filepath}")
                except Exception as e:
                    self.show_error_message(f"Export failed: {str(e)}")

        elif fmt == 'pdf':
            self._export_pdf(tab_name, headers, rows, period_lbl)

        elif fmt == 'print':
            self._print_report(tab_name, headers, rows, period_lbl)

    def _export_master_workbook(self):
        """Export all 4 sections into a single master Excel workbook with multiple tabs"""
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter

        wb = Workbook()
        # Remove default sheet
        wb.remove(wb.active)

        period_lbl = self.cached_results.get('period_label', 'Custom Period')
        company = get_setting('company_name', 'AC Service Billing') or 'AC Service Billing'

        sections = [
            ('Profit & Loss', self.pnl_table),
            ('GST Report', self.gst_table),
            ('Parts Usage', self.parts_table),
            ('Team Performance', self.tech_table),
        ]

        header_fill = PatternFill('solid', fgColor='0F172A')
        header_font = Font(name='Segoe UI', size=10, bold=True, color='FFFFFF')
        zebra_fill = PatternFill('solid', fgColor='F8FAFC')
        total_fill = PatternFill('solid', fgColor='EFF6FF')
        total_font = Font(name='Segoe UI', size=10, bold=True, color='0F172A')
        thin_border = Border(
            left=Side(style='thin', color='E2E8F0'), right=Side(style='thin', color='E2E8F0'),
            top=Side(style='thin', color='E2E8F0'), bottom=Side(style='thin', color='E2E8F0')
        )

        for sheet_title, table in sections:
            ws = wb.create_sheet(title=sheet_title[:31])
            ws.views.sheetView[0].showGridLines = True

            # Title block
            ws.cell(row=1, column=1, value=f"{company} - {sheet_title}").font = Font(name='Segoe UI', size=14, bold=True, color='0F172A')
            ws.cell(row=2, column=1, value=f"Period: {period_lbl} | Generated: {datetime.now().strftime('%d-%m-%Y %H:%M')}").font = Font(name='Segoe UI', size=9, color='64748B')

            col_count = table.columnCount()
            row_count = table.rowCount()

            # Headers at Row 4
            for c in range(col_count):
                cell = ws.cell(row=4, column=c + 1, value=table.horizontalHeaderItem(c).text())
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal='center', vertical='center')

            # Data rows
            for r in range(row_count):
                excel_row = r + 5
                is_total = (r == row_count - 1)  # last row is totals
                for c in range(col_count):
                    item = table.item(r, c)
                    val_str = item.text() if item else ''
                    cell = ws.cell(row=excel_row, column=c + 1)

                    # Try to parse currency / number
                    clean_str = val_str.replace(self.currency_symbol, '').replace(',', '').strip()
                    is_pct = clean_str.endswith('%')
                    if is_pct:
                        clean_str = clean_str[:-1]

                    try:
                        num = float(clean_str)
                        if is_pct:
                            cell.value = num / 100.0
                            cell.number_format = '0.0%'
                            cell.alignment = Alignment(horizontal='center', vertical='center')
                        elif '.' in clean_str:
                            cell.value = num
                            cell.number_format = f'"{self.currency_symbol}"#,##0.00'
                            cell.alignment = Alignment(horizontal='right', vertical='center')
                        else:
                            cell.value = int(num)
                            if self.currency_symbol in val_str:
                                cell.number_format = f'"{self.currency_symbol}"#,##0'
                                cell.alignment = Alignment(horizontal='right', vertical='center')
                            else:
                                cell.number_format = '#,##0'
                                cell.alignment = Alignment(horizontal='center', vertical='center')
                    except ValueError:
                        cell.value = val_str
                        cell.alignment = Alignment(horizontal='left' if c == 0 else 'center', vertical='center')

                    cell.border = thin_border
                    if is_total:
                        cell.fill = total_fill
                        cell.font = total_font
                    elif r % 2 == 1:
                        cell.fill = zebra_fill
                        cell.font = Font(name='Segoe UI', size=9.5, color='1E293B')
                    else:
                        cell.font = Font(name='Segoe UI', size=9.5, color='1E293B')

            # Auto-fit columns
            for col in ws.columns:
                max_len = max(len(str(c.value or '')) for c in col)
                col_letter = get_column_letter(col[0].column)
                ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

        filepath, _ = QFileDialog.getSaveFileName(
            self, 'Save Executive Financial Master Workbook',
            f"financial_master_report_{period_lbl.replace(' ', '_').lower()}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx",
            'Excel Files (*.xlsx)'
        )
        if filepath:
            try:
                wb.save(filepath)
                self.show_success_message(f"Complete Master Workbook Exported:\n{filepath}")
            except Exception as e:
                self.show_error_message(f"Export failed: {str(e)}")

    def _export_pdf(self, tab_name, headers, rows, period_lbl):
        try:
            from reportlab.lib.pagesizes import A4, landscape
            from reportlab.lib import colors
            from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
            from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
            from reportlab.lib.units import mm
        except ImportError:
            self.show_warning_message("ReportLab library is required for PDF export.")
            return

        company = get_setting('company_name', 'AC Service Billing') or 'AC Service Billing'

        filepath, _ = QFileDialog.getSaveFileName(
            self, 'Save PDF Report',
            f"{tab_name.replace(' ', '_')}_{period_lbl.replace(' ', '_')}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pdf",
            'PDF Files (*.pdf)'
        )
        if not filepath:
            return

        doc = SimpleDocTemplate(filepath, pagesize=landscape(A4),
                                rightMargin=15 * mm, leftMargin=15 * mm,
                                topMargin=12 * mm, bottomMargin=12 * mm)
        styles = getSampleStyleSheet()
        elements = []

        title_style = ParagraphStyle('CustomTitle', parent=styles['Title'],
                                     fontSize=15, textColor=colors.HexColor('#0F172A'),
                                     spaceAfter=2, alignment=0)
        sub_style = ParagraphStyle('CustomSub', parent=styles['Normal'],
                                   fontSize=8.5, textColor=colors.HexColor('#64748B'),
                                   spaceAfter=10, alignment=0)

        elements.append(Paragraph(f"{company} - {tab_name} Report", title_style))
        elements.append(Paragraph(f"Period: {period_lbl} | Generated: {datetime.now().strftime('%d-%m-%Y %H:%M')}", sub_style))

        styled_headers = [
            Paragraph(h, ParagraphStyle('Hdr', fontSize=7.5, textColor=colors.white, fontName='Helvetica-Bold', alignment=1))
            for h in headers
        ]
        data_rows = [styled_headers]
        for r_idx, row in enumerate(rows):
            is_total = (r_idx == len(rows) - 1)
            row_cells = []
            for cell in row:
                f_name = 'Helvetica-Bold' if is_total else 'Helvetica'
                t_color = colors.HexColor('#0F172A') if is_total else colors.HexColor('#1E293B')
                row_cells.append(Paragraph(str(cell), ParagraphStyle('C', fontSize=7, textColor=t_color, fontName=f_name, alignment=1)))
            data_rows.append(row_cells)

        col_w = [doc.width / len(headers)] * len(headers)
        t = Table(data_rows, colWidths=col_w, repeatRows=1)
        t.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#0F172A')),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
            ('ROWBACKGROUNDS', (0, 1), (-1, -2), [colors.white, colors.HexColor('#F8FAFC')]),
            ('BACKGROUND', (0, -1), (-1, -1), colors.HexColor('#EFF6FF')),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ]))
        elements.append(t)

        try:
            doc.build(elements)
            self.show_success_message(f"PDF report saved:\n{filepath}")
        except Exception as e:
            self.show_error_message(f"Failed to generate PDF: {str(e)}")

    def _print_report(self, tab_name, headers, rows, period_lbl):
        from PySide6.QtPrintSupport import QPrinter, QPrintDialog
        from PySide6.QtGui import QTextDocument

        company = get_setting('company_name', 'AC Service Billing') or 'AC Service Billing'

        html = f"""
        <html>
        <head><style>
            body {{ font-family: 'Segoe UI', Arial, sans-serif; padding: 15px; color: #0f172a; }}
            h1 {{ color: #0f172a; font-size: 16pt; margin-bottom: 2px; }}
            .subtitle {{ color: #64748b; font-size: 8.5pt; margin-bottom: 15px; }}
            table {{ width: 100%; border-collapse: collapse; font-size: 8pt; }}
            th {{ background-color: #0f172a; color: white; padding: 6px; text-align: center; border: 1px solid #0f172a; }}
            td {{ padding: 5px 6px; border: 1px solid #cbd5e1; text-align: center; }}
            tr:nth-child(even) {{ background-color: #f8fafc; }}
            tr:last-child {{ background-color: #eff6ff; font-weight: bold; }}
        </style></head>
        <body>
        <h1>{company} - {tab_name} Report</h1>
        <p class="subtitle">Period: {period_lbl} | Generated: {datetime.now().strftime('%d-%m-%Y %H:%M')}</p>
        <table>
        <thead><tr>{''.join(f'<th>{h}</th>' for h in headers)}</tr></thead>
        <tbody>
        {''.join(f'<tr>{"".join(f"<td>{c}</td>" for c in row)}</tr>' for row in rows)}
        </tbody>
        </table>
        </body></html>
        """
        doc = QTextDocument()
        doc.setHtml(html)
        printer = QPrinter(QPrinter.PrinterMode.HighResolution)
        dialog = QPrintDialog(printer, self)
        if dialog.exec() == QPrintDialog.DialogCode.Accepted:
            doc.print_(printer)
